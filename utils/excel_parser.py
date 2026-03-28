import io
import re
import logging
from datetime import datetime, date
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.utils.datetime import from_excel

logger = logging.getLogger(__name__)

# -----------------------------
# Normalization / value helpers
# -----------------------------
def normalize_text(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    s = str(val).replace("\xa0", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def empty_to_blank(val: Any) -> Any:
    if val is None:
        return ""
    if isinstance(val, str):
        val = re.sub(r"\s+", " ", val).strip()
        if val.upper() in {"", "N/A", "NA"}:
            return ""
    return val


JUNK_PATTERNS = [
    r"CONTINUE ON SEPARATE",
    r"CS FORM 212",
    r"WORK EXPERIENCE",
    r"LEARNING AND DEVELOPMENT",
    r"OTHER INFORMATION",
    r"DATE OF BIRTH",
    r"INCLUSIVE DATES",
    r"DEPARTMENT / AGENCY",
    r"STATUS OF APPOINTMENT",
    r"GOV'?T SERVICE",
    r"NAME OF CHILDREN",
    r"WRITE FULL",
    r"TYPE OF L&D",
    r"MANAGERIAL",
    r"SUPERVISORY",
    r"BASIC EDUCATION",
    r"HIGHEST LEVEL",
    r"SCHOLARSHIP",
    r"INCLUDE PRIVATE EMPLOYMENT",
    r"START FROM YOUR RECENT",
    r"DESCRIPTION OF DUTIES",
    r"DO NOT ABBREVIATE",
    r"MEMBERSHIP IN ASSOCIATION",
    r"SIGNATURE",
    r"DATE ACCOMPLISHED",
    r"PHOTO",
]


def is_junk(val: Any) -> bool:
    s = normalize_text(val).upper()
    if not s:
        return True
    if re.fullmatch(r"\d+\.", s):
        return True
    if s in {"DATE", "FROM", "TO"}:
        return True
    return any(re.search(p, s, re.IGNORECASE) for p in JUNK_PATTERNS)


def parse_date(val: Any, epoch=None) -> str:
    """
    Converts true Excel dates / full date strings to YYYY-MM-DD.
    Keeps partial dates like 'JUN,2001' as-is.
    """
    val = empty_to_blank(val)
    if val == "":
        return ""

    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")

    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")

    if isinstance(val, (int, float)):
        try:
            return from_excel(val, epoch=epoch).strftime("%Y-%m-%d")
        except Exception:
            return str(val)

    s = normalize_text(val)
    if not s:
        return ""

    if s.upper() == "PRESENT":
        return "PRESENT"

    # Only parse FULL dates, not partial month/year text
    full_date_formats = [
        "%d/%m/%Y",
        "%d/%m/%y",
        "%m/%d/%Y",
        "%Y-%m-%d",
        "%B %d, %Y",
        "%B %d,%Y",
        "%b %d, %Y",
        "%b %d,%Y",
    ]

    for fmt in full_date_formats:
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass

    return s


# -----------------------------
# Merged-cell aware readers
# -----------------------------
def build_merged_lookup(ws) -> Dict[tuple, Any]:
    lookup = {}
    for rng in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = rng.bounds
        top_left_val = ws.cell(min_row, min_col).value
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                lookup[(r, c)] = top_left_val
    return lookup


def get_raw(ws, row: int, col: int, merged_lookup=None) -> Any:
    if not ws or row < 1 or col < 1 or row > ws.max_row or col > ws.max_column:
        return None
    val = ws.cell(row=row, column=col).value
    if val is None and merged_lookup is not None:
        val = merged_lookup.get((row, col))
    return val


def get_val(ws, row: int, col: int, merged_lookup=None) -> Any:
    return empty_to_blank(get_raw(ws, row, col, merged_lookup))


# -----------------------------
# Finders / parsers
# -----------------------------
def find_sheet(wb, target: str):
    # exact first
    for name in wb.sheetnames:
        if name.strip().upper() == target.upper():
            return wb[name]

    # fuzzy fallback
    pattern = re.compile(rf"(^|[^A-Z0-9]){re.escape(target)}([^A-Z0-9]|$)", re.IGNORECASE)
    for name in wb.sheetnames:
        if pattern.search(name):
            return wb[name]

    return None


def find_row_any(ws, start_row: int, end_row: int, pattern: str, cols=None, merged_lookup=None) -> int:
    if not ws:
        return -1

    regex = re.compile(pattern, re.IGNORECASE)
    cols = cols or range(1, ws.max_column + 1)

    for r in range(start_row, min(end_row, ws.max_row) + 1):
        for c in cols:
            val = get_raw(ws, r, c, merged_lookup)
            if isinstance(val, str) and regex.search(val):
                return r
    return -1


def first_nonempty_in_rows(ws, rows, col: int, merged_lookup=None) -> str:
    for r in rows:
        v = get_val(ws, r, col, merged_lookup)
        if v != "":
            return v
    return ""


def parse_table(
    ws,
    start_row: int,
    end_row: int,
    mappings: Dict[str, Any],
    required_keys=None,
    merged_lookup=None,
    row_filter=None,
) -> List[Dict[str, Any]]:
    if not ws:
        return []

    if isinstance(required_keys, str):
        required_keys = [required_keys]
    required_keys = required_keys or []

    result = []

    for r in range(start_row, min(end_row, ws.max_row) + 1):
        row_obj = {}
        non_empty_count = 0

        for key, cfg in mappings.items():
            if isinstance(cfg, int):
                col = cfg
                transform = None
            else:
                col = cfg["col"]
                transform = cfg.get("transform")

            raw = get_val(ws, r, col, merged_lookup)
            val = transform(raw) if (transform and raw != "") else raw

            row_obj[key] = val
            if val not in ("", None):
                non_empty_count += 1

        if non_empty_count == 0:
            continue

        joined = " ".join(normalize_text(v) for v in row_obj.values() if v not in ("", None))
        if joined and is_junk(joined):
            continue

        if required_keys and not any(row_obj.get(k) not in ("", None) for k in required_keys):
            continue

        if row_filter and not row_filter(row_obj):
            continue

        result.append(row_obj)

    return result


def parse_list(ws, start_row: int, end_row: int, col: int, merged_lookup=None) -> List[str]:
    if not ws:
        return []

    res = []
    for r in range(start_row, min(end_row, ws.max_row) + 1):
        val = get_val(ws, r, col, merged_lookup)
        if val != "" and not is_junk(val):
            res.append(val)
    return res


# -----------------------------
# Main parser
# -----------------------------
def parse_pds_excel(file_content: bytes) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
    logger.info("Workbook sheets found: %s", wb.sheetnames)

    c1 = find_sheet(wb, "C1") or (wb.worksheets[0] if wb.worksheets else None)
    c2 = find_sheet(wb, "C2") or (wb.worksheets[1] if len(wb.worksheets) > 1 else c1)
    c3 = find_sheet(wb, "C3") or (wb.worksheets[2] if len(wb.worksheets) > 2 else c1)
    c4 = find_sheet(wb, "C4") or (wb.worksheets[3] if len(wb.worksheets) > 3 else None)

    if not c1:
        raise ValueError("No readable worksheet found.")

    m1 = build_merged_lookup(c1)
    m2 = build_merged_lookup(c2) if c2 else {}
    m3 = build_merged_lookup(c3) if c3 else {}
    m4 = build_merged_lookup(c4) if c4 else {}

    # Section anchors
    family_row = find_row_any(c1, 1, 100, r"II\.\s*FAMILY BACKGROUND", merged_lookup=m1)
    education_row = find_row_any(c1, 1, 150, r"III\.\s*EDUCATIONAL BACKGROUND", merged_lookup=m1)

    # Find Father/Mother within family section more robustly
    father_row = find_row_any(c1, family_row, family_row + 20, r"FATHER'?S (?:NAME|SURNAME)", merged_lookup=m1) if family_row > 0 else -1
    mother_row = find_row_any(c1, family_row, family_row + 25, r"MOTHER['S]?\\s+MAIDEN\\s+NAME", merged_lookup=m1) if family_row > 0 else -1
    if mother_row < 0 and family_row > 0:
        mother_row = find_row_any(c1, family_row, family_row + 25, r"MOTHER[’]?\\s+MAIDEN\\s+NAME", merged_lookup=m1)
    if mother_row < 0:
        mother_row = 46

    civil_service_row = find_row_any(c2, 1, 50, r"IV\.\s*CIVIL SERVICE ELIGIBILITY", merged_lookup=m2)
    work_row = find_row_any(c2, 1, 80, r"V\.\s*WORK EXPERIENCE", merged_lookup=m2)

    voluntary_row = find_row_any(c3, 1, 50, r"VI\.\s*VOLUNTARY WORK", merged_lookup=m3)
    training_row = find_row_any(c3, 1, 80, r"VII\.\s*LEARNING AND DEVELOPMENT", merged_lookup=m3)
    other_row = find_row_any(c3, 1, 100, r"VIII\.\s*OTHER INFORMATION", merged_lookup=m3)

    civil_status = first_nonempty_in_rows(c1, range(10, 20), 16, m1)

    # Dynamic starts based on observed layouts
    education_data_start = education_row + 4 if education_row > 0 else 54
    civil_service_data_start = civil_service_row + 3 if civil_service_row > 0 else 5
    work_data_start = work_row + 5 if work_row > 0 else 18
    voluntary_data_start = voluntary_row + 4 if voluntary_row > 0 else 6
    training_data_start = training_row + 4 if training_row > 0 else 18
    other_data_start = other_row + 2 if other_row > 0 else 42

    data = {
        "personal_info": {
            "surname": get_val(c1, 10, 4, m1),
            "first_name": get_val(c1, 11, 4, m1),
            "middle_name": get_val(c1, 12, 4, m1),
            "name_extension": re.sub(
                r"(?i)^NAME EXTENSION \(JR\., SR\)\s*", "",
                normalize_text(get_raw(c1, 11, 12, m1))
            ).replace("N/A", "").strip(),
            "date_of_birth": parse_date(get_raw(c1, 13, 4, m1), wb.epoch),
            "place_of_birth": get_val(c1, 15, 4, m1),
            "sex_at_birth": get_val(c1, 16, 4, m1),
            "civil_status": civil_status,
            "height": get_val(c1, 22, 4, m1),
            "weight": get_val(c1, 24, 4, m1),
            "blood_type": get_val(c1, 25, 4, m1),

            # Revised 2025 form
            "umid_no": get_val(c1, 27, 4, m1),
            "pagibig_no": get_val(c1, 29, 4, m1),
            "philhealth_no": get_val(c1, 31, 4, m1),
            "philsys_no": get_val(c1, 32, 4, m1),
            "tin_no": get_val(c1, 33, 4, m1),
            "agency_employee_no": get_val(c1, 34, 4, m1),

            # backward-compatible alias if your downstream code still expects gsis_no
            "gsis_no": get_val(c1, 27, 4, m1),

            "telephone_no": get_val(c1, 32, 9, m1),
            "mobile_no": get_val(c1, 33, 9, m1),
            "email_address": get_val(c1, 34, 9, m1),

            "residential_address": {
                "house_no": get_val(c1, 17, 9, m1),
                "street": get_val(c1, 17, 12, m1),
                "subdivision": get_val(c1, 19, 9, m1),
                "barangay": get_val(c1, 19, 12, m1),
                "city": get_val(c1, 22, 9, m1),
                "province": get_val(c1, 22, 12, m1),
                "zip_code": get_val(c1, 24, 9, m1),
            },
            "permanent_address": {
                "house_no": get_val(c1, 25, 9, m1),
                "street": get_val(c1, 25, 12, m1),
                "subdivision": get_val(c1, 27, 9, m1),
                "barangay": get_val(c1, 27, 12, m1),
                "city": get_val(c1, 29, 10, m1),
                "province": get_val(c1, 29, 13, m1),
                "zip_code": get_val(c1, 31, 9, m1),
            },
        },

        "family_background": {
            "spouse": {
                "surname": get_val(c1, 36, 4, m1),
                "first_name": get_val(c1, 37, 4, m1),
                "middle_name": get_val(c1, 38, 4, m1),
                "occupation": get_val(c1, 39, 4, m1),
                "employer_name": get_val(c1, 40, 4, m1),
                "business_address": get_val(c1, 41, 4, m1),
                "telephone_no": get_val(c1, 42, 4, m1),
            },
        "father": {
            "surname": get_val(c1, father_row, 4, m1) if father_row > 0 else "",
            "first_name": get_val(c1, father_row + 1, 4, m1) if father_row > 0 else "",
            "middle_name": get_val(c1, father_row + 2, 4, m1) if father_row > 0 else "",
        },
        "mother": {
            # Values start one row below the header line:
            # row +1 = surname, +2 = first name, +3 = middle name.
            "maiden_surname": get_val(c1, mother_row + 1, 4, m1) if mother_row > 0 else "",
            "first_name": get_val(c1, mother_row + 2, 4, m1) if mother_row > 0 else "",
            "middle_name": get_val(c1, mother_row + 3, 4, m1) if mother_row > 0 else "",
        },
        "children": parse_table(
            c1,
            37,
            (education_row - 1 if education_row > 0 else 49),
            {
                    "name": 9,
                    "date_of_birth": {"col": 13, "transform": lambda v: parse_date(v, wb.epoch)},
                },
                required_keys="name",
                merged_lookup=m1,
                row_filter=lambda row: row.get("name") not in ("", None),
            ),
        },

        "educational_background": parse_table(
            c1,
            education_data_start,
            58,
            {
                "level": 2,
                "school_name": 4,
                "degree": 7,
                "from": 10,
                "to": 11,
                "highest_level": 12,
                "year_graduated": 13,
                "scholarships": 14,
            },
            required_keys=["school_name", "degree"],
            merged_lookup=m1,
            row_filter=lambda row: (
                row.get("school_name") not in ("", None) or
                row.get("degree") not in ("", None)
            ),
        ),

        "civil_service_eligibility": parse_table(
            c2,
            civil_service_data_start,
            (work_row - 1 if work_row > 0 else 12),
            {
                "eligibility": 1,
                "rating": 6,
                "exam_date": {"col": 7, "transform": lambda v: parse_date(v, wb.epoch)},
                "exam_place": 9,
                "license_number": 10,
                "license_valid_until": {"col": 11, "transform": lambda v: parse_date(v, wb.epoch)},
            },
            required_keys="eligibility",
            merged_lookup=m2,
            row_filter=lambda row: row.get("eligibility") not in ("", None),
        ),

        # Revised 2025 form no longer matches your old salary/pay-grade mapping
        "work_experience": parse_table(
            c2,
            work_data_start,
            45,
            {
                "from": {"col": 1, "transform": lambda v: parse_date(v, wb.epoch)},
                "to": {"col": 3, "transform": lambda v: parse_date(v, wb.epoch)},
                "position": 4,
                "department": 7,
                "appointment_status": 10,
                "govt_service": 11,
            },
            required_keys=["position", "department"],
            merged_lookup=m2,
            row_filter=lambda row: row.get("position") not in ("", None),
        ),

        "voluntary_work": parse_table(
            c3,
            voluntary_data_start,
            (training_row - 1 if training_row > 0 else 13),
            {
                "organization": 1,
                "from": {"col": 5, "transform": lambda v: parse_date(v, wb.epoch)},
                "to": {"col": 6, "transform": lambda v: parse_date(v, wb.epoch)},
                "hours": 7,
                "position": 8,
            },
            required_keys="organization",
            merged_lookup=m3,
            row_filter=lambda row: row.get("organization") not in ("", None),
        ),

        "training_programs": parse_table(
            c3,
            training_data_start,
            (other_row - 1 if other_row > 0 else 39),
            {
                "title": 1,
                "from": {"col": 5, "transform": lambda v: parse_date(v, wb.epoch)},
                "to": {"col": 6, "transform": lambda v: parse_date(v, wb.epoch)},
                "hours": 7,
                "type": 8,
                "conducted_by": 9,
            },
            required_keys="title",
            merged_lookup=m3,
            row_filter=lambda row: row.get("title") not in ("", None),
        ),

        "other_information": {
            "skills": parse_list(c3, other_data_start, 48, 1, m3),
            "distinctions": parse_list(c3, other_data_start, 48, 3, m3),
            "memberships": parse_list(c3, other_data_start, 48, 9, m3),
        },
    }

    # -------------------------------------------------------
    # Fix key name: parser used 'license_valid_until' but the
    # rest of the codebase (autoFillForm2, db_to_pds_prefill)
    # expects 'license_date'. Normalise here.
    # -------------------------------------------------------
    for elig in data["civil_service_eligibility"]:
        if "license_valid_until" in elig:
            elig["license_date"] = elig.pop("license_valid_until")

    # -------------------------------------------------------
    # Page 4 (Sheet C4) — References, Gov't ID, Declarations
    # -------------------------------------------------------
    if c4:
        # ---- References (Item 41) -------------------------
        ref_row = find_row_any(c4, 1, 60, r"REFERENCES?.*\(Person not", merged_lookup=m4)
        if ref_row < 0: 
            ref_row = find_row_any(c4, 1, 60, r"REFERENCES", merged_lookup=m4)

        ref_data_start = ref_row + 2 if ref_row > 0 else 4
        data["references"] = parse_table(
            c4,
            ref_data_start,
            ref_data_start + 2,  # Read exactly 3 rows (e.g. 52, 53, 54)
            {
                "name":    1,
                "address": 5,
                "contact": 9,
            },
            required_keys="name",
            merged_lookup=m4,
            row_filter=lambda row: row.get("name") not in ("", None),
        )

        # ---- Government-Issued ID (Item 42) ---------------
        gov_id_row = find_row_any(c4, 1, 80, r"Government Issued ID", merged_lookup=m4)
        gov_id_data_row = gov_id_row + 2 if gov_id_row > 0 else -1
        data["government_id"] = {
            "type":                  get_val(c4, gov_id_data_row, 1, m4)  if gov_id_data_row > 0 else "",
            "number":                get_val(c4, gov_id_data_row, 5, m4)  if gov_id_data_row > 0 else "",
            "date_place_of_issuance":get_val(c4, gov_id_data_row, 9, m4)  if gov_id_data_row > 0 else "",
        }

        # ---- Declarations (Items 34-40) -------------------
        # Locate each question label and read answer + detail
        def _decl_read(ws, merged, question_pattern, start=1, stop=100):
            """Find a declaration question row and return (yes_no, detail)."""
            r = find_row_any(ws, start, stop, question_pattern, merged_lookup=merged)
            if r < 0:
                return "", ""
            # Yes/No is typically in col 12-13
            yn_raw = get_val(ws, r, 12, merged) or get_val(ws, r, 13, merged) or ""
            yn = "yes" if str(yn_raw).strip().upper() in ("YES", "Y", "X", "✓", "✔") else ""
            # Because detail box placement is highly variable and often spans multiple rows
            # below the question, we leave detail string blank to avoid grabbing question text.
            return yn, ""

        decl_start = 1
        decl_stop  = gov_id_row - 1 if gov_id_row > 0 else 80

        yn34a, det34a = _decl_read(c4, m4, r"within the third degree",  decl_start, decl_stop)
        yn34b, det34b = _decl_read(c4, m4, r"within the fourth degree", decl_start, decl_stop)
        yn35a, det35a = _decl_read(c4, m4, r"guilty of any administrative offense", decl_start, decl_stop)
        yn35b, det35b = _decl_read(c4, m4, r"criminally charged before any court",  decl_start, decl_stop)
        yn36,  det36  = _decl_read(c4, m4, r"convicted of any crime",               decl_start, decl_stop)
        yn37,  det37  = _decl_read(c4, m4, r"separated from the service",           decl_start, decl_stop)
        yn38a, det38a = _decl_read(c4, m4, r"candidate in a national or local",     decl_start, decl_stop)
        yn38b, det38b = _decl_read(c4, m4, r"resigned from the government service", decl_start, decl_stop)
        yn39,  det39  = _decl_read(c4, m4, r"immigrant or permanent resident",      decl_start, decl_stop)
        yn40a, det40a = _decl_read(c4, m4, r"member of any indigenous group",       decl_start, decl_stop)
        yn40b, det40b = _decl_read(c4, m4, r"person with disability",               decl_start, decl_stop)
        yn40c, det40c = _decl_read(c4, m4, r"solo parent",                          decl_start, decl_stop)

        data["declarations"] = {
            # Yes/No answers (stored for form radio pre-selection)
            "q34a": yn34a, "q34b": yn34b,
            "q35a": yn35a, "q35b": yn35b,
            "q36":  yn36,  "q37":  yn37,
            "q38a": yn38a, "q38b": yn38b,
            "q39":  yn39,
            "q40a": yn40a, "q40b": yn40b, "q40c": yn40c,
            # Detail fields (matched to DB column names autoFillForm4 reads)
            "related_appointing_authority_3rd_degree": det34a,
            "related_appointing_authority_4th_degree": det34b,
            "administrative_offense_details":          det35a,
            "criminal_charge_details":                 det35b,
            "conviction_details":                     det36,
            "separation_details":                     det37,
            "election_candidacy_details":             det38a,
            "resignation_campaign_details":           det38b,
            "immigrant_status_country":               det39,
            "indigenous_group":                       det40a,
            "pwd_id":                                 det40b,
            "solo_parent_id":                         det40c,
        }
    else:
        # Sheet C4 not found — leave declarations/references/gov_id empty
        data["declarations"] = {}
        data["references"]   = []
        data["government_id"] = {
            "type": "", "number": "", "date_place_of_issuance": ""
        }

    return data