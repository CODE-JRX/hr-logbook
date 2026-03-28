import openpyxl, os

files = [os.path.join('static/uploads/temp_pds', f) for f in os.listdir('static/uploads/temp_pds') if f.endswith('.xlsx')]
latest_file = max(files, key=os.path.getmtime)

wb = openpyxl.load_workbook(latest_file, data_only=True)
c4 = wb['C4'] if 'C4' in wb.sheetnames else list(wb.worksheets)[-1]
m4 = {}
for rng in c4.merged_cells.ranges:
    min_col, min_row, max_col, max_row = rng.bounds
    top_left_val = c4.cell(min_row, min_col).value
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            m4[(r, c)] = top_left_val

with open('test_output.txt', 'w', encoding='utf-8') as f:
    f.write(f'Testing file: {latest_file}\n')
    for r in range(1, 100):
        for c in range(1, 15):
            val = str(c4.cell(r, c).value or '').strip()
            if not val:
                val = str(m4.get((r, c)) or '').strip()
            
            v_upper = val.upper()
            if 'REFEREN' in v_upper or '41' in v_upper or '34' in v_upper or '35' in v_upper or '36' in v_upper or '37' in v_upper or '38' in v_upper or '39' in v_upper or '40' in v_upper or 'GOVERNMENT' in v_upper:
                f.write(f'Row {r}, Col {c}: {repr(val)}\n')
