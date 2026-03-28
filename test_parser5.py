import openpyxl, os, sys, json

files = [os.path.join('static/uploads/temp_pds', f) for f in os.listdir('static/uploads/temp_pds') if f.endswith('.xlsx')]
latest_file = max(files, key=os.path.getmtime)

wb = openpyxl.load_workbook(latest_file, data_only=True)
c4 = wb['C4'] if 'C4' in wb.sheetnames else list(wb.worksheets)[-1]
m4 = {}
for rng in c4.merged_cells.ranges:
    min_col, min_row, max_col, max_row = rng.bounds
    top_left_val = c4.cell(min_row, min_col).value
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            m4[(r, c)] = top_left_val

def get_v(r, c):
    v = c4.cell(r, c).value
    if v is None: v = m4.get((r, c))
    return str(v or '').strip()

with open('parser_dump5.txt', 'w', encoding='utf-8') as f:
    f.write("--- REFS ---\n")
    for r in range(50, 65):
        f.write(f"Row {r}: 1='{get_v(r, 1)}', 5='{get_v(r, 5)}', 9='{get_v(r, 9)}'\n")
        
    f.write("---\n")
    for r in range(1, 45):
        if 'YES' in get_v(r, 5) or 'YES' in get_v(r, 4) or 'YES' in get_v(r, 6) or 'YES' in get_v(r, 7):
            f.write(f"Row {r}: 4='{get_v(r, 4)}', 5='{get_v(r, 5)}'\n")
        elif 'YES' in get_v(r, 12).upper() or 'NO' in get_v(r, 12).upper():
            f.write(f"Checkbox Row {r}: 12='{get_v(r, 12)}', 13='{get_v(r, 13)}'\n")
