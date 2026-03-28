import openpyxl, os

files = [os.path.join('static/uploads/temp_pds', f) for f in os.listdir('static/uploads/temp_pds') if f.endswith('.xlsx')]
latest_file = max(files, key=os.path.getmtime)

wb = openpyxl.load_workbook(latest_file, data_only=True)
c4 = wb['C4'] if 'C4' in wb.sheetnames else list(wb.worksheets)[-1]
m4 = {}
for rng in c4.merged_cells.ranges:
    min_col, min_row, max_col, max_row = rng.bounds
    top_left_val = c4.cell(min_row, min_col).value
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            m4[(r, c)] = top_left_val

with open('test_output2.txt', 'w', encoding='utf-8') as f:
    for r in range(1, 50):
        row_vals = []
        for c in range(1, 15):
            val = str(c4.cell(r, c).value or '').strip()
            if not val:
                val = str(m4.get((r, c)) or '').strip()
            row_vals.append(val)
        if any(row_vals):
            f.write(f'Row {r}: {" | ".join(row_vals)}\n')
